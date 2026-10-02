"""Semantic chunk boundaries and API/CLI equivalence regressions."""

import re
from types import SimpleNamespace

import pytest
from docx import Document as WordDocument
from langchain.schema import Document

from application.ports.document_parser import FileMeta, SplitContext
from application.services.document_loader import S3DocumentLoader
from application.services.document_content import DocumentContentProcessor
from domain.domain_profile.profiles.decree import DecreeDomainProfile
from domain.domain_profile.profiles.legal import LegalDomainProfile
from domain.domain_profile.registry import DomainProfileRegistry
from domain.entities.raw_document import RawDocument
from domain.value_objects.doc_domain import DocDomain
from domain.value_objects.page_content_type import PageContentType
from infrastructure.domain_profile.settings_adapter import DomainSettingsAdapter
from infrastructure.ingestion.document_parser_adapter import (
    IngestionDocumentParser,
    IngestionDocumentSplitter,
)
from infrastructure.ml.ingestion.splitting import split_documents
from infrastructure.ml.ingestion.markdown import extract_date_from_filename
from infrastructure.ml.langchain_document_parser import LangchainDocumentParser, LangchainDocumentSplitter


@pytest.fixture
def pipeline():
    settings = DomainSettingsAdapter()
    registry = DomainProfileRegistry()
    for profile_type in (DecreeDomainProfile, LegalDomainProfile):
        profile = profile_type(settings)
        registry.register(profile)
        for default in profile.config_defaults():
            settings.set(default.key, profile.key, default.value)
        settings.set(f"{profile.key}_chunk_size", profile.key, "1200")
        settings.set(f"{profile.key}_chunk_overlap", profile.key, "200")
    return LangchainDocumentParser(registry, settings), LangchainDocumentSplitter(registry, settings)


def normalize(text):
    return re.sub(r"\s+", " ", text).strip()


def test_instruction_steps_that_fit_are_kept_whole(pipeline):
    _, splitter = pipeline
    step = "3. Заполните основные поля\n" + "Проверьте сведения о товаре. " * 11
    step += "После заполнения нажмите: Создать, Подписать с ЭЦП, Отправить."
    assert len(step) < 550
    text = "Инструкция по созданию акта\n1. Откройте документ\n" + "Описание. " * 25 + "\n" + step
    chunks = splitter.split([RawDocument(text, {"source": "guide.txt"})])
    assert any(normalize(step) in normalize(c.page_content) for c in chunks)
    assert all(len(c.page_content) <= 550 for c in chunks)


def test_large_step_repeats_heading_without_exceeding_budget(pipeline):
    _, splitter = pipeline
    text = "1. Откройте документ\nОткройте карточку.\n2. Заполните сведения\n" + "Проверьте сведения. " * 80
    chunks = splitter.split([RawDocument(text, {"source": "guide.txt"})])
    continuations = [c for c in chunks if c.metadata.get("step_number") == "2"]
    assert len(continuations) > 1
    assert all("2. Заполните сведения" in c.page_content for c in continuations)
    assert all(len(c.page_content) <= 550 for c in chunks)


def test_short_neighboring_steps_are_packed_with_all_numbers(pipeline):
    _, splitter = pipeline
    text = (
        "Guide\n1. Open document\nOpen the card.\n"
        "2. Fill data\nFill the required fields.\n3. Send\nSign and send."
    )
    chunks = splitter.split([RawDocument(text, {"source": "guide.txt"})])
    assert len(chunks) == 1
    assert chunks[0].page_content == text
    assert chunks[0].metadata["step_numbers"] == ["1", "2", "3"]


def test_subpoints_keep_parent_condition_and_distinct_numbers(pipeline):
    _, splitter = pipeline
    condition = "21. Только для электронных заявлений действуют следующие правила:"
    text = condition + "\n21.1. Срок подачи — пять дней.\n21.2. Подавать через личный кабинет."
    chunks = splitter.split([RawDocument(text, {"source": "rules.txt"})], DocDomain.DECREE.value)
    subpoints = [c for c in chunks if c.metadata.get("unit_kind") == "subpoint_num"]
    assert [c.metadata["subpoint_num_number"] for c in subpoints] == ["21.1", "21.2"]
    assert all(condition in c.page_content for c in subpoints)
    assert all(c.metadata["point_number"] == "21" for c in subpoints)
    assert "21.2." not in subpoints[0].page_content
    assert "21.1." not in subpoints[1].page_content


def test_overflow_keeps_article_marker_and_bounded_final_context(pipeline):
    _, splitter = pipeline
    text = "Статья 15. Порядок подачи\nПорядок применяется к заявлениям. "
    text += "Заявитель представляет необходимые документы в установленном порядке " * 35 + "."
    chunks = splitter.split([RawDocument(text, {"source": "law.txt"})], DocDomain.LEGAL.value)
    assert len(chunks) > 1
    assert all(len(c.page_content) <= 1200 for c in chunks)
    assert any("Статья 15." in c.page_content for c in chunks)
    assert all(c.metadata.get("article_number") == "15." for c in chunks)
    assert not any(c.page_content.strip() in ("15", "15.") for c in chunks)


def test_tables_keep_original_position_and_repeated_header():
    rows = "\n".join(f"| item-{i} | {i} |" for i in range(20))
    docs = [
        Document(page_content="Before.", metadata={"source": "guide.md"}),
        Document(
            page_content="| Name | Fee |\n|---|---|\n" + rows,
            metadata={"source": "guide.md", "content_type": PageContentType.TABLE.value},
        ),
        Document(page_content="After.", metadata={"source": "guide.md"}),
    ]
    chunks = split_documents(docs)
    assert chunks[0].page_content == "Before."
    assert chunks[-1].page_content == "After."
    assert len(chunks[1:-1]) == 2
    assert all(c.page_content.startswith("| Name | Fee |") for c in chunks[1:-1])
    assert [c.metadata["chunk_index"] for c in chunks] == [1, 2, 3, 4]


def test_pdf_cross_page_sentence_and_chunk_page_ranges(pipeline):
    _, splitter = pipeline
    docs = [
        RawDocument("Service is available only if", {"source": "guide.pdf", "page": 1}),
        RawDocument("the application is submitted before April 30.", {"source": "guide.pdf", "page": 2}),
    ]
    chunks = splitter.split(docs)
    assert len(chunks) == 1
    assert (
        normalize(chunks[0].page_content)
        == "Service is available only if the application is submitted before April 30."
    )
    assert chunks[0].metadata["pages"] == [1, 2]
    assert "page" not in chunks[0].metadata
    long_docs = [
        RawDocument("Page one sentence. " * 40, {"source": "long.pdf", "page": 1}),
        RawDocument("Page two sentence. " * 40, {"source": "long.pdf", "page": 2}),
    ]
    chunks = splitter.split(long_docs)
    assert chunks[0].metadata["pages"] == [1]
    assert chunks[-1].metadata["pages"] == [2]
    assert all("_page_spans" not in c.metadata for c in chunks)


def test_pdf_table_stays_a_barrier_between_text_pages(pipeline):
    _, splitter = pipeline
    docs = [
        RawDocument("Before.", {"source": "guide.pdf", "page": 1}),
        RawDocument(
            "| Name | Fee |\n|---|---|\n| A | 5 |",
            {"source": "guide.pdf", "page": 1, "content_type": PageContentType.TABLE.value},
        ),
        RawDocument("After.", {"source": "guide.pdf", "page": 2}),
    ]
    chunks = splitter.split(docs)
    assert [c.metadata.get("content_type") for c in chunks] == [None, PageContentType.TABLE.value, None]


def test_repeated_pdf_text_still_advances_to_the_correct_page(pipeline):
    _, splitter = pipeline
    text = "Same sentence repeated. " * 80
    docs = [RawDocument(text, {"source": "repeat.pdf", "page": page}) for page in (1, 2)]
    chunks = splitter.split(docs)
    assert chunks[0].metadata["pages"] == [1]
    assert chunks[-1].metadata["pages"] == [2]
    page_one = [c.page_content for c in chunks if c.metadata["pages"] == [1]]
    page_two = [c.page_content for c in chunks if c.metadata["pages"] == [2]]
    assert page_one == page_two


@pytest.mark.parametrize("extension", [".md", ".docx", ".rtf", ".txt", ".pdf"])
def test_api_cli_same_chunks_and_structural_metadata(tmp_path, pipeline, extension, monkeypatch):
    parser, splitter = pipeline
    path = tmp_path / ("download-123" + extension)
    if extension == ".docx":
        word = WordDocument()
        word.add_heading("Guide", 1)
        word.add_paragraph("1. Open document\nOpen the document card.")
        word.add_paragraph("2. Fill data\nFill all required fields and send the document.")
        table = word.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Fee"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "5"
        word.save(path)
    elif extension == ".rtf":
        text = "Заявитель представляет документы в установленном порядке " * 40
        path.write_text(
            r"{\rtf1\ansi УКАЗ № 5 от 1 января 2026 г.\par ПОСТАНОВЛЯЮ:\par 1. " + text + r"\par}",
            encoding="utf-8",
        )
    elif extension == ".pdf":
        path.write_bytes(b"fixture")
        monkeypatch.setattr(
            "infrastructure.ml.langchain_document_parser.parse_pdf",
            lambda _: [
                Document(
                    page_content="Service is available only if", metadata={"source": str(path), "page": 1}
                ),
                Document(
                    page_content="the application is submitted before April 30.",
                    metadata={"source": str(path), "page": 2},
                ),
            ],
        )
    else:
        path.write_text(
            "# Guide\n\n1. Open document\nOpen the document card.\n"
            "2. Fill data\nFill all fields.\n\n| Name | Fee |\n|---|---|\n| A | 5 |",
            encoding="utf-8",
        )
    filename = "guide_2026-01-15" + extension
    api_docs = parser.parse(path)
    for doc in api_docs:
        doc.metadata["source"] = filename
    domain = DocDomain.DECREE.value if extension == ".rtf" else DocDomain.GENERAL.value
    api_chunks = splitter.split(api_docs, domain)
    meta = FileMeta("s3://bucket/" + filename, filename, extension, path.stat().st_size)
    cli_docs = IngestionDocumentParser(parser).parse(path, meta)
    cli_chunks = IngestionDocumentSplitter(splitter).split(cli_docs, SplitContext(domain=domain))
    assert [c.page_content for c in api_chunks] == [c.page_content for c in cli_chunks]
    transport = {"source", "filename", "extension", "size_bytes"}
    assert [{k: v for k, v in c.metadata.items() if k not in transport} for c in api_chunks] == [
        {k: v for k, v in c.metadata.items() if k not in transport} for c in cli_chunks
    ]
    assert all(c.metadata["doc_date"] == "2026-01-15" for c in cli_chunks)
    if extension == ".rtf":
        assert all(len(c.page_content) <= 1200 for c in cli_chunks)
        assert not any(c.page_content.strip() == "1." for c in cli_chunks)


def test_batch_loader_uses_each_source_domain_before_splitting(pipeline):
    _, splitter = pipeline
    loader = S3DocumentLoader(None, None, IngestionDocumentSplitter(splitter), SimpleNamespace())
    docs = [
        RawDocument("Ordinary text.", {"source": "plain.txt"}),
        RawDocument("1. Первое правило.\n2. Второе правило.", {"source": "law.txt"}),
    ]
    chunks = loader.split_docs(
        docs, source_domains={"plain.txt": DocDomain.GENERAL.value, "law.txt": DocDomain.DECREE.value}
    )
    legal = [c for c in chunks if c.metadata["source"] == "law.txt"]
    assert len(legal) == 2
    assert all(c.metadata["unit_kind"] == "point" for c in legal)
    assert all(c.metadata["doc_domain"] == DocDomain.DECREE.value for c in legal)


def test_empty_table_with_section_does_not_become_a_prefix_only_chunk():
    docs = [
        Document(
            page_content=" \n", metadata={"section": "Fees", "content_type": PageContentType.TABLE.value}
        )
    ]
    assert split_documents(docs) == []


def test_chapterless_article_keeps_identity_on_overflow_chunks(pipeline):
    _, splitter = pipeline
    text = "Статья 15. Порядок подачи\n" + "Заявитель передаёт документы в установленном порядке. " * 25
    assert len(text) < 1500
    chunks = splitter.split([RawDocument(text, {"source": "law.txt"})], DocDomain.LEGAL.value)
    assert len(chunks) > 1
    assert all(c.metadata["article_number"] == "15." for c in chunks)
    assert all("article 15." in c.metadata["section"] for c in chunks)


def test_docx_generic_footnote_part_is_parsed(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    word = WordDocument()
    word.add_heading("Guide", 1)
    word.add_paragraph("The document body contains a complete instruction.")
    xml = (
        b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:footnote w:id="1"><w:p><w:r><w:t>Important exception.</w:t></w:r></w:p></w:footnote>'
        b'</w:footnotes>'
    )
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        xml,
        word.part.package,
    )
    word.part.relate_to(part, RELATIONSHIP_TYPE.FOOTNOTES)
    path = tmp_path / "notes.docx"
    word.save(path)
    docs = LangchainDocumentParser().parse(path)
    assert any("Important exception." in doc.page_content for doc in docs)


@pytest.mark.asyncio
async def test_api_content_processor_and_cli_construct_identical_final_chunks(tmp_path, pipeline):
    parser, splitter = pipeline
    path = tmp_path / "download.md"
    path.write_text(
        "# Guide\n\n## Registration\n\n" + "Complete a registration form. " * 35, encoding="utf-8"
    )
    filename = "guide_2026-01-15.md"
    extractor = SimpleNamespace(extract_date_from_filename=extract_date_from_filename)
    processor = DocumentContentProcessor(parser, splitter, extractor, None, None, None)
    ctx = SimpleNamespace(
        docs=parser.parse(path), original_filename=filename, doc_domain=DocDomain.GENERAL.value
    )
    api_chunks = await processor.split(ctx)
    meta = FileMeta("s3://bucket/" + filename, filename, ".md", path.stat().st_size)
    cli_chunks = IngestionDocumentSplitter(splitter).split(
        IngestionDocumentParser(parser).parse(path, meta), SplitContext()
    )
    assert [c.page_content for c in api_chunks] == [c.page_content for c in cli_chunks]
    assert all(c.page_content.count("[Раздел:") == 1 for c in api_chunks)
    assert all(len(c.page_content) <= 550 for c in api_chunks)
    assert all(c.metadata["char_count"] == len(c.page_content) for c in api_chunks)
