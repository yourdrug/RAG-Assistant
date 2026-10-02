"""Public parser/splitter contracts captured before chunking changes."""

import pytest
from docx import Document as WordDocument
from langchain.schema import Document

from application.ports.document_parser import FileMeta, SplitContext
from domain.entities.raw_document import RawDocument
from domain.value_objects.doc_domain import DocDomain
from domain.value_objects.page_content_type import PageContentType
from infrastructure.ingestion.document_parser_adapter import (
    IngestionDocumentParser,
    IngestionDocumentSplitter,
)
from infrastructure.ml.ingestion import parse_docx, parse_docx_sections, split_documents
from infrastructure.ml.langchain_document_parser import LangchainDocumentParser, LangchainDocumentSplitter


def test_markdown_parser_keeps_breadcrumbs_and_table_headers(tmp_path):
    path = tmp_path / "guide.md"
    path.write_text(
        "# Guide\n\n## Fees\n\nPrices below.\n\n| Name | Fee |\n|---|---|\n| A | 5 |", encoding="utf-8"
    )
    docs = LangchainDocumentParser().parse(path)
    assert [d.metadata["section"] for d in docs] == ["Guide > Fees", "Guide > Fees"]
    assert docs[1].metadata["content_type"] == PageContentType.TABLE.value
    assert "| Name | Fee |" in docs[1].page_content
    assert all(d.metadata["doc_title"] == "Guide" for d in docs)


def test_docx_public_parsers_keep_body_order_and_breadcrumbs(tmp_path):
    path = tmp_path / "guide.docx"
    word = WordDocument()
    word.add_heading("Guide", 1)
    word.add_heading("Fees", 2)
    word.add_paragraph("Before the table.")
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Fee"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "5"
    word.add_paragraph("After the table.")
    word.save(path)
    text, _ = parse_docx(path)
    assert text.index("Before") < text.index("| Name") < text.index("After")
    sections = parse_docx_sections(path)
    assert len(sections) == 3
    assert all(heading == "Guide > Fees" for heading, _ in sections)
    assert sections[1][1].startswith("\x00TABLE:")


@pytest.mark.parametrize("extension", [".txt", ".md", ".docx", ".rtf", ".pdf"])
def test_cli_parser_supported_formats(extension):
    assert IngestionDocumentParser().supports(extension)
    assert not IngestionDocumentParser().supports(".exe")


def test_cli_parser_keeps_storage_identity(tmp_path):
    path = tmp_path / "guide.txt"
    path.write_text("Document content with enough meaningful text.", encoding="utf-8")
    meta = FileMeta("s3://bucket/guide.txt", "guide.txt", ".txt", path.stat().st_size)
    docs = IngestionDocumentParser().parse(path, meta)
    assert len(docs) == 1
    assert docs[0].metadata["source"] == meta.source_key
    assert docs[0].metadata["filename"] == meta.filename
    assert docs[0].page_content == path.read_text(encoding="utf-8")


@pytest.mark.parametrize("splitter", [LangchainDocumentSplitter(), IngestionDocumentSplitter()])
def test_splitter_retains_text_and_custom_metadata(splitter):
    text = "A short paragraph with a complete thought."
    docs = [RawDocument(text, {"source": "guide.txt", "custom": "kept"})]
    if isinstance(splitter, LangchainDocumentSplitter):
        chunks = splitter.split(docs, domain=DocDomain.GENERAL.value)
    else:
        chunks = splitter.split(docs, SplitContext())
    assert len(chunks) == 1
    assert chunks[0].page_content == text
    assert chunks[0].metadata["custom"] == "kept"


def test_table_batches_repeat_header_and_preserve_all_rows():
    rows = [f"| item-{i} | {i} |" for i in range(31)]
    table = "| Name | Fee |\n|---|---|\n" + "\n".join(rows)
    chunks = split_documents(
        [Document(page_content=table, metadata={"content_type": PageContentType.TABLE.value})]
    )
    assert len(chunks) == 3
    assert all(c.page_content.startswith("| Name | Fee |\n|---|---|") for c in chunks)
    assert [row for c in chunks for row in c.page_content.splitlines()[2:]] == rows


def test_split_documents_empty_input_and_final_positions():
    assert split_documents([]) == []
    docs = [Document(page_content="word " * 200, metadata={"source": "guide.txt"})]
    chunks = split_documents(docs)
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(1, len(chunks) + 1))
    assert all(c.metadata["total_chunks"] == len(chunks) for c in chunks)
    assert all(c.metadata["char_count"] == len(c.page_content) for c in chunks)
