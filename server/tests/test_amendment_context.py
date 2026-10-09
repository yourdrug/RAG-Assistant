"""Replacement evidence survives parsing, chunking, persistence and budgeting."""

import docx
import pytest
from langchain.schema import Document

from domain.value_objects.chunk_context import extract_chunk_context
from domain.value_objects.page_content_type import PageContentType
from infrastructure.ml.ingestion.amendment_context import attach_table_replacement_scope
from infrastructure.ml.ingestion.table_chunks import table_document_chunks
from infrastructure.ml.langchain_document_parser import LangchainDocumentParser
from infrastructure.ml.rag.evidence_focus import (
    amendment_anchors,
    format_generation_messages,
    reading_anchors,
)
from infrastructure.ml.rag.rag_formatting import format_docs_with_selection
from infrastructure.ml.rag.rag_prompts import build_prompt

FIELD = "Дополнительный четырехзначный код по классификатору дополнительной таможенной информации"
TARGET = "в таблице 4.1 - структура и формат сообщения BLRWBL (ЭТТН) позицию"
BRIDGE = "заменить позицией:"
OLD = "С"
NEW = "М - для товаров, подлежащих прослеживаемости, С - для прочих товаров"
QUESTION = f"Как изменился статус обязательности поля «{FIELD}» (позиция 57 таблицы 4.1)?"


@pytest.fixture
def replacement_docs(tmp_path):
    path = tmp_path / "amendment.docx"
    source = docx.Document()
    source.add_paragraph(TARGET)
    for status in (OLD, NEW):
        table = source.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ('"57', FIELD + " an..4", status), strict=True):
            cell.text = value
        if status == OLD:
            source.add_paragraph(BRIDGE)
    source.save(path)
    return LangchainDocumentParser().parse(path)


def table_parts(docs):
    return [doc for doc in docs if doc.metadata.get("content_type") == PageContentType.TABLE.value]


def test_parser_preserves_original_replacement_on_both_values(replacement_docs):
    old, new = table_parts(replacement_docs)
    expected = "\n".join([TARGET, old.page_content, BRIDGE, new.page_content])
    for part in (old, new):
        assert part.metadata["parent_units"] == [{"content": expected}]
        assert extract_chunk_context(part.metadata)["parent_units"][0]["content"] == expected


@pytest.mark.parametrize("selected_table", [0, 1])
def test_either_retrieved_table_carries_complete_atomic_scope(replacement_docs, selected_table):
    raw = table_parts(replacement_docs)[selected_table]
    chunks = table_document_chunks(Document(page_content=raw.page_content, metadata=raw.metadata), 250, 5)
    assert all(len(chunk.page_content) <= 250 for chunk in chunks)
    # Simulate the persisted interpretation fields restored onto a search hit.
    hit = Document(page_content=chunks[0].page_content, metadata=extract_chunk_context(chunks[0].metadata))
    context, selected = format_docs_with_selection([hit], 5000, context_counter=len)
    scope = raw.metadata["parent_units"][0]["content"]
    assert scope in context
    assert selected
    assert context.index(TARGET) < context.index(BRIDGE) < context.index(NEW)
    assert f"| {OLD} |" in context
    assert amendment_anchors(context, QUESTION) == [context]
    messages = format_generation_messages(
        build_prompt(question=QUESTION), context=context, history=[], question=QUESTION
    )
    assert scope in messages[-2].content
    # Insufficient room must exclude the whole scope, never just its connector.
    too_small, selected = format_docs_with_selection([hit], len(context) - 1, context_counter=len)
    assert too_small == ""
    assert selected == []


@pytest.mark.parametrize("mutation", ["missing_operation", "different_heading", "different_source"])
def test_cannot_infer_replacement_from_disconnected_rows(replacement_docs, mutation):
    for doc in replacement_docs:
        doc.metadata.pop("parent_units", None)
    if mutation == "missing_operation":
        replacement_docs[2].page_content = "Дополнительные сведения"
    elif mutation == "different_heading":
        replacement_docs[-1].metadata["section"] = "Другая поправка"
    else:
        replacement_docs[-1].metadata["source"] = "Другой акт"
    attach_table_replacement_scope(replacement_docs)
    assert all("parent_units" not in doc.metadata for doc in replacement_docs)


def test_focus_requires_explicit_operation_and_requested_table(replacement_docs):
    disconnected = "\n".join(doc.page_content for doc in table_parts(replacement_docs))
    assert amendment_anchors(disconnected, QUESTION) == []
    assert reading_anchors(disconnected, QUESTION) == []
    scope = table_parts(replacement_docs)[0].metadata["parent_units"][0]["content"]
    assert amendment_anchors(scope.replace("таблице 4.1", "таблице 4.2"), QUESTION) == []
    assert amendment_anchors(scope, QUESTION.replace("четырех", "четырёх")) == [scope]
