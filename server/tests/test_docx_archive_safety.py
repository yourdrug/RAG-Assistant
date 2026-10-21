"""DOCX archive limits apply before upload and before either worker parser."""

import struct
from io import BytesIO
from zipfile import ZIP_BZIP2, ZIP_DEFLATED, ZIP_STORED, ZipFile

import docx
import pytest
from fastapi.exceptions import RequestValidationError
from pydantic import ValidationError

from infrastructure.ml.ingestion.docx import parse_docx, parse_docx_sections
from presentation.api.schemas.validators import FileContent, validate_uploaded_file


def make_archive(entries, *, compression=ZIP_STORED):
    output = BytesIO()
    with ZipFile(output, "w", compression=compression) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return output.getvalue()


def archive_with_declared_sizes(sizes):
    """Forge central-directory sizes to test large limits without allocating them."""
    data = bytearray(make_archive([(f"word/part{i}.xml", b"x") for i in range(len(sizes))]))
    start = 0
    for size in sizes:
        start = data.index(b"PK\x01\x02", start)
        struct.pack_into("<II", data, start + 20, size, size)
        start += 46
    return bytes(data)


@pytest.mark.parametrize(
    "data,reason",
    [
        (archive_with_declared_sizes([101 * 1024 * 1024]), "uncompressed"),
        (archive_with_declared_sizes([55 * 1024 * 1024] * 2), "uncompressed"),
        (make_archive([(f"word/{i}.xml", b"x") for i in range(1001)]), "entries"),
        (make_archive([("word/document.xml", b"x" * 512 * 1024)], compression=ZIP_DEFLATED), "ratio"),
    ],
    ids=["single_large_entry", "total_large_entries", "too_many_entries", "compression_bomb"],
)
def test_structural_validation_rejects_unsafe_docx_archives(data, reason):
    with pytest.raises(ValidationError, match=reason):
        FileContent(data=data, filename="unsafe.docx", check_structural=True)


def test_upload_rejects_zip_bomb_with_public_422_error():
    data = make_archive([("word/document.xml", b"x" * 512 * 1024)], compression=ZIP_DEFLATED)
    with pytest.raises(RequestValidationError) as error:
        validate_uploaded_file(data, "unsafe.docx")
    assert error.value.errors()[0]["loc"] == ("body", "file")


def test_underreported_expanded_size_cannot_bypass_docx_limits():
    data = bytearray(make_archive([("word/document.xml", b"x" * 512 * 1024)], compression=ZIP_DEFLATED))
    central_directory = data.index(b"PK\x01\x02")
    struct.pack_into("<I", data, central_directory + 24, 512)
    with pytest.raises(RequestValidationError):
        validate_uploaded_file(bytes(data), "forged.docx")


def test_unbounded_compression_methods_are_rejected_before_parsing():
    data = make_archive([("word/document.xml", b"document text")], compression=ZIP_BZIP2)
    with pytest.raises(RequestValidationError):
        validate_uploaded_file(data, "bzip.docx")


@pytest.mark.parametrize("parser", [parse_docx, parse_docx_sections])
def test_worker_rejects_unsafe_docx_even_without_api_upload_validation(tmp_path, parser):
    path = tmp_path / "unsafe.docx"
    document = docx.Document()
    document.add_paragraph("Normal body text")
    document.save(path)
    with ZipFile(path, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/unused.xml", b"x" * 512 * 1024)
    with pytest.raises(ValueError, match="ratio"):
        parser(path)


def test_normal_docx_passes_upload_and_both_worker_parsers(tmp_path):
    path = tmp_path / "safe.docx"
    document = docx.Document()
    document.add_heading("Safe document", level=1)
    document.add_paragraph("Normal body text")
    document.save(path)
    validated = validate_uploaded_file(path.read_bytes(), path.name)
    assert validated.data == path.read_bytes()
    assert "Normal body text" in parse_docx(path)[0]
    assert any("Normal body text" in text for text in (s.content for s in parse_docx_sections(path)))
