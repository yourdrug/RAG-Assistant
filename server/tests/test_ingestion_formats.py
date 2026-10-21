"""Contracts for format dispatch and the shared ingestion representation."""

from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock

from docx import Document as WordDocument
from docx.enum.text import WD_BREAK
import fitz  # type: ignore[import-untyped]
from langchain.schema import Document
import pytest

from config import settings
from domain.ingestion.preview import DryRunPageResult, PDFQualityReport
from domain.knowledge.enums import PageContentType
from infrastructure.ingestion.document_parser_adapter import IngestionDocumentParser
from infrastructure.ml.extraction.extraction_adapter import MLPDFQualityAssessor
from infrastructure.ml.ingestion import docx, ocr, rtf
from infrastructure.ml.ingestion.registry import FORMAT_HANDLERS, FormatContext
from infrastructure.ml.ingestion.sections import ParsedFormat, Section
from infrastructure.ml.langchain_document_parser import LangchainDocumentParser
from infrastructure.ml.preview import docx_preview_strategy
from infrastructure.ml.preview.factory import PreviewStrategyFactory


def test_docx_reads_archive_and_renders_each_table_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "guide.docx"
    word = WordDocument()
    word.core_properties.title = "Core title"
    word.add_heading("Guide", 1)
    word.add_paragraph("Before the table.")
    word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    table = word.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).text = "Value"
    word.add_paragraph("After the table.")
    word.save(str(path))
    read = Mock(wraps=WordDocument)
    render = Mock(wraps=docx.docx_table_to_markdown)
    monkeypatch.setattr("docx.Document", read)
    monkeypatch.setattr(docx, "docx_table_to_markdown", render)

    docs = LangchainDocumentParser().parse(path)

    read.assert_called_once_with(str(path))
    assert render.call_count == 1
    assert [d.metadata.get("content_type") for d in docs] == [None, "table", None]
    assert [d.metadata["section"] for d in docs] == ["Guide"] * 3
    assert all(d.metadata["doc_title"] == "Core title" for d in docs)
    assert all(d.metadata["pages"] == [1, 2] and d.metadata["table_count"] == 1 for d in docs)
    assert docs[0].page_content == "Before the table."
    assert docs[-1].page_content == "After the table."


def test_ingestion_and_preview_dispatch_through_the_same_handler(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    parse = Mock(return_value=ParsedFormat([Section("Heading", "Document body.")]))
    strategy = Mock()
    preview = Mock(return_value=strategy)
    handler = replace(FORMAT_HANDLERS[".md"], extensions=(".custom",), parse=parse, preview=preview)
    monkeypatch.setitem(FORMAT_HANDLERS, ".custom", handler)
    assert IngestionDocumentParser().supports(".CUSTOM")
    docs = LangchainDocumentParser().parse(tmp_path / "guide.CUSTOM")
    assert docs[0].metadata["section"] == "Heading"
    assert PreviewStrategyFactory.for_extension(".CUSTOM") is strategy
    assert ".custom" in PreviewStrategyFactory.supported_extensions()
    parse.assert_called_once_with(tmp_path / "guide.CUSTOM", FormatContext())
    preview.assert_called_once_with(FormatContext())


def test_literal_old_marker_remains_text(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    text = "\x00TABLE: This is literal text, not a table."
    path.write_text(text, encoding="utf-8")
    docs = LangchainDocumentParser().parse(path)
    assert docs[0].page_content == text
    assert docs[0].metadata.get("content_type") != PageContentType.TABLE


def test_pdf_quality_uses_domain_report(tmp_path: Path) -> None:
    path = tmp_path / "guide.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.new_page()
        pdf.save(path)
    report = MLPDFQualityAssessor().assess(
        path, [Document(page_content="Meaningful words on this page. " * 20, metadata={"page": 1})]
    )
    assert isinstance(report, PDFQualityReport)
    assert report.total_pages == 2
    assert report.n_missing == 1
    assert report.n_ok == 1
    assert report.is_low_quality


@pytest.mark.parametrize("paddle_text", ["", "Paddle recognized text"])
def test_pdf_auto_falls_back_to_surya_only_for_empty_paddle_result(
    paddle_text: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "ocr_engine", "auto")
    paddle = Mock(return_value=paddle_text)
    surya = Mock(return_value="Surya recognized text")
    monkeypatch.setattr(ocr, "ocr_image_paddle", paddle)
    monkeypatch.setattr(ocr, "ocr_image_surya", surya)
    with fitz.open() as pdf:
        pdf.new_page(width=20, height=20)
        results = ocr.ocr_pdf_pages(pdf, [1], "guide.pdf")
    assert results == {1: paddle_text or "Surya recognized text"}
    assert paddle.call_count == 1
    assert surya.call_count == (0 if paddle_text else 1)


def test_docx_auto_keeps_paddle_only_policy(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # Image decoding is exercised by VL tests; here the contract is engine selection.
    image = Mock()
    monkeypatch.setattr(docx_preview_strategy, "_extract_images_without_alt", Mock(return_value=[b"image"]))
    monkeypatch.setattr("PIL.Image.open", Mock(return_value=image))
    monkeypatch.setattr(settings, "ocr_engine", "auto")
    paddle = Mock(return_value="")
    surya = Mock()
    monkeypatch.setattr(docx_preview_strategy, "ocr_image_paddle", paddle)
    monkeypatch.setattr(docx_preview_strategy, "ocr_image_surya", surya)
    unit = DryRunPageResult(page=1, type=PageContentType.IMAGE_ONLY)
    units, _, _ = docx_preview_strategy.DocxPreviewStrategy().ocr_problem_units(
        tmp_path / "guide.docx", [unit], [1]
    )
    paddle.assert_called_once_with(image)
    surya.assert_not_called()
    assert units == [unit]


def test_rtf_structural_failure_preserves_flat_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "guide.rtf"
    path.write_text(
        r"{\rtf1\ansi Meaningful document text with enough words for ingestion.}", encoding="ascii"
    )
    monkeypatch.setattr(rtf, "parse_rtf_sections", Mock(side_effect=ValueError("invalid structure")))
    docs = LangchainDocumentParser().parse(path)
    assert docs[0].page_content == rtf.parse_rtf(path)[0]
    assert "falling back to flat parse" in caplog.text


@pytest.mark.parametrize("extension", [".doc", ".DOC"])
def test_legacy_doc_is_not_accepted_by_ingestion(tmp_path: Path, extension: str) -> None:
    assert not IngestionDocumentParser().supports(extension)
    assert extension.lower() not in settings.supported_extensions
    with pytest.raises(RuntimeError, match="Unsupported format"):
        LangchainDocumentParser().parse(tmp_path / f"legacy{extension}")
