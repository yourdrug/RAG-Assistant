"""Regression scenarios for mixed PDF pages, instruction scope and prompt budgets."""

from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import fitz
import pytest
from docx import Document as WordDocument
from langchain.schema import Document

from domain.domain_profile.profiles.decree import DecreeDomainProfile
from domain.domain_profile.registry import DomainProfileRegistry
from domain.entities.raw_document import RawDocument
from domain.exceptions import ContextBudgetExceededError
from domain.value_objects.doc_domain import DocDomain
from domain.value_objects.chat_context import ChatContext
from domain.value_objects.llm_provider import Breadth
from domain.value_objects.page_content_type import PageContentType
from domain.value_objects.roles import UserKind, UserRole
from domain.value_objects.stream_events import StatusEvent
from domain.value_objects.user_context import UserContext
from infrastructure.domain_profile.settings_adapter import DomainSettingsAdapter
from infrastructure.ml.ingestion.pdf import parse_pdf, pymupdf_table_to_markdown
from infrastructure.ml.ingestion.splitting import split_documents
from infrastructure.ml.langchain_document_parser import LangchainDocumentParser, LangchainDocumentSplitter
from infrastructure.ml.rag.rag_formatting import CHARS_PER_TOKEN, format_docs, format_docs_with_selection
from infrastructure.ml.rag.rag_steps import step_build_context
from infrastructure.ml.rag_pipeline import RagPipelineState


@pytest.fixture
def pipeline():
    settings = DomainSettingsAdapter()
    registry = DomainProfileRegistry()
    profile = DecreeDomainProfile(settings)
    registry.register(profile)
    for default in profile.config_defaults():
        settings.set(default.key, profile.key, default.value)
    return LangchainDocumentParser(registry, settings), LangchainDocumentSplitter(registry, settings)


def test_mixed_pdf_keeps_text_before_between_and_after_tables(tmp_path):
    path = tmp_path / "mixed.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((60, 90), "Before: fees apply only to new applications.", fontsize=10)
        page.insert_text((300, 115), "Before-right: additional eligibility.", fontsize=10)
        for top in (150, 300):
            for y in (top, top + 30, top + 60):
                page.draw_line((60, y), (440, y))
            for x in (60, 250, 440):
                page.draw_line((x, top), (x, top + 60))
            for x, y, value in (
                (70, top + 20, "Service"),
                (260, top + 20, "Fee"),
                (70, top + 50, f"Item-{top}"),
                (260, top + 50, "50"),
            ):
                page.insert_text((x, y), value, fontsize=10)
        page.insert_text((60, 260), "Between: the second table covers additional services.", fontsize=10)
        page.insert_text((60, 400), "After: existing customers are exempt from payment.", fontsize=10)
        assert len(page.find_tables().tables) == 2
        pdf.save(path)
    docs = parse_pdf(path)
    chunks = split_documents(docs)
    content = "\n".join(c.page_content for c in chunks)
    assert content.index("Before:") < content.index("Item-150") < content.index("Between:")
    assert content.index("Before-right:") < content.index("Item-150")
    assert content.index("Between:") < content.index("Item-300") < content.index("After:")
    assert content.count("Item-150") == content.count("Item-300") == 1
    assert all(c.metadata["page"] == 1 for c in chunks)
    assert all(not any(key.startswith("_table") for key in c.metadata) for c in chunks)


def test_multiline_docx_cells_remain_one_logical_row(pipeline, tmp_path):
    parser, splitter = pipeline
    word = WordDocument()
    table = word.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Service\ncategory"
    table.cell(0, 1).text = "Fee"
    table.cell(1, 0).text = "Registration"
    table.cell(1, 1).text = "50\nTax included"
    path = tmp_path / "multiline.docx"
    word.save(path)
    chunks = splitter.split(parser.parse(path))
    assert len(chunks) == 1
    assert chunks[0].metadata["table_row_start"] == chunks[0].metadata["table_row_end"] == 1
    lines = chunks[0].page_content.splitlines()
    assert len(lines) == 3
    assert "Service" in lines[0] and "category" in lines[0]
    assert "Registration" in lines[2] and "50" in lines[2] and "Tax included" in lines[2]
    assert chunks[0].metadata["table_header"] == "\n".join(lines[:2])


def test_multiline_pdf_cells_remain_one_logical_row():
    table = SimpleNamespace(extract=lambda: [["Service\ncategory", "Fee"], ["A", "50\nTax included"]])
    chunks = split_documents(
        [
            Document(
                page_content=pymupdf_table_to_markdown(table),
                metadata={"content_type": PageContentType.TABLE.value},
            )
        ]
    )
    assert len(chunks) == 1
    assert chunks[0].metadata["table_row_end"] == 1
    assert len(chunks[0].page_content.splitlines()) == 3


def test_each_instruction_step_retains_the_complete_intro(pipeline):
    _, splitter = pipeline
    condition = "Only existing customers may use this procedure."
    intro = condition + "\n" + "Review the eligibility conditions before continuing. " * 8
    text = intro + "\n1. Send application\n" + "Fill in all required fields. " * 11
    text += "\n2. Pay fee\n" + "Pay the fee using a bank transfer. " * 9
    chunks = splitter.split([RawDocument(text, {"source": "guide.txt"})])
    steps = [c for c in chunks if c.metadata.get("step_number")]
    assert steps
    assert all(condition in format_docs([c]) for c in steps)
    assert all(c.metadata["parent_units"][-1]["content"] == intro.strip() for c in steps)
    assert all(len(c.page_content) <= 550 for c in chunks)


def test_packed_steps_keep_intro_even_when_it_is_in_a_separate_chunk(pipeline):
    _, splitter = pipeline
    intro = "Use these steps only after approval. " * 16
    text = intro + "\n1. Open\nOpen the card.\n2. Send\nSend the application."
    chunks = splitter.split([RawDocument(text, {"source": "guide.txt"})])
    packed = next(c for c in chunks if c.metadata.get("step_numbers") == ["1", "2"])
    assert intro.strip() in format_docs([packed])


def test_oversized_parent_is_never_truncated_or_allowed_to_exceed_budget(pipeline):
    _, splitter = pipeline
    intro = "21. " + "Eligibility conditions apply to applications. " * 700
    text = intro + "\n21.1. Submit the passport.\n21.2. Submit the form."
    child = next(
        c
        for c in splitter.split([RawDocument(text, {"source": "rules.txt"})], DocDomain.DECREE.value)
        if c.metadata.get("subpoint_num_number") == "21.2"
    )
    context, selected = format_docs_with_selection([child], max_context_tokens=6000)
    assert len(context) <= 6000 * CHARS_PER_TOKEN
    assert selected == []
    assert context == ""


def test_oversized_document_does_not_discard_later_usable_sources():
    oversized = Document(page_content="large " * 100, metadata={"source": "large.txt"})
    small = Document(page_content="Complete statement.", metadata={"source": "small.txt"})
    context, selected = format_docs_with_selection([(oversized, 0.9), (small, 0.8)], max_context_tokens=30)
    assert len(context) <= 30 * CHARS_PER_TOKEN
    assert selected == [(small, 0.8)]
    assert "[1] small.txt" in context and "[2]" not in context


def test_budget_includes_separators():
    docs = [Document(page_content="x" * 14, metadata={"source": "a"}) for _ in range(2)]
    context, selected = format_docs_with_selection(docs, max_context_tokens=10)
    assert len(context) <= 10 * CHARS_PER_TOKEN
    assert len(selected) == 1


@pytest.mark.parametrize("tokens", [0, -1])
def test_nonpositive_context_budget_is_rejected(tokens):
    with pytest.raises(ValueError, match="positive"):
        format_docs_with_selection([], max_context_tokens=tokens)


def test_no_usable_context_stops_before_generation():
    from test_rag_pipeline import _make_rag

    state = RagPipelineState(
        rag=_make_rag(),
        t_pipeline_start=0,
        question="Explain the eligibility rules.",
        ctx=ChatContext(
            user_id=12, user_kind=UserKind.INTERNAL.value, user_role=UserRole.USER.value, user_group_ids=[]
        ),
        user=UserContext(12, UserKind.INTERNAL, UserRole.USER),
        access_filter=None,
        retrieval_filter=None,
        breadth=Breadth.NARROW,
        docs=[(Document(page_content="Complete conditions. " * 1500, metadata={"source": "rules.pdf"}), 0.9)],
    )
    with pytest.raises(ContextBudgetExceededError):
        step_build_context(state, None)
    assert state._prompt_docs == []


@pytest.mark.asyncio
async def test_stream_reports_context_budget_error_without_answer_or_sources():
    from presentation.api.routes.chat import chat_stream
    from presentation.api.schemas import ChatRequest

    async def stream_chat(*args, **kwargs):
        yield StatusEvent(stage="searching")
        raise ContextBudgetExceededError()

    response = await chat_stream(
        ChatRequest(question="Укажите порядок подачи заявления."),
        SimpleNamespace(is_disconnected=AsyncMock(return_value=False)),
        current_user=SimpleNamespace(id=12, kind=UserKind.INTERNAL.value, role=UserRole.USER.value),
        chat_service=SimpleNamespace(stream_chat=stream_chat),
        log=MagicMock(),
    )
    events = "".join([event async for event in response.body_iterator])
    assert "event: error" in events and ContextBudgetExceededError.code in events
    assert "Уточните вопрос" in events
    assert '"text"' not in events and "event: done" not in events
