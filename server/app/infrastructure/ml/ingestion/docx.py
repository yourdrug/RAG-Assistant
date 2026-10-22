"""DOCX parsing — python-docx based text/table/section extraction."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from domain.knowledge.enums import PageContentType
from infrastructure.ml.ingestion.sections import Section
from infrastructure.ml.ingestion.table_formatting import markdown_cell, render_table
from infrastructure.storage.docx_validation import validate_docx_archive


def _has_page_break(paragraph) -> bool:
    """Check if a paragraph contains a manual page break."""
    for run in paragraph.runs:
        for br in run._element.findall(qn("w:br")):
            if br.get(qn("w:type")) == "page":
                return True

    return False


def _paragraph_full_text(paragraph) -> str:
    """Full paragraph text, including runs nested inside hyperlinks."""
    t_tag = qn("w:t")
    tab_tag = qn("w:tab")
    br_tag = qn("w:br")
    parts: list[str] = []

    for node in paragraph._element.iter():
        if node.tag == t_tag:
            if node.text:
                parts.append(node.text)
        elif node.tag == tab_tag:
            parts.append("\t")
        elif node.tag == br_tag:
            parts.append("\n")

    return "".join(parts)


# --- Numbering format resolution — delegated to _docx_numbering ---

from infrastructure.ml.ingestion.docx_numbering import (  # noqa: E402, F401
    _get_numbering_formats,
    ListNumberer,
    _paragraph_list_info,
    _to_roman,
)


# --- Extraction helpers -----------------------------------------------------


def _extract_headers_footers(doc) -> dict:
    """Collect unique header/footer text across all sections."""
    headers: list[str] = []
    footers: list[str] = []
    seen_headers: set[str] = set()
    seen_footers: set[str] = set()

    for section in doc.sections:
        for container, seen, bucket in (
            (section.header, seen_headers, headers),
            (section.footer, seen_footers, footers),
        ):
            if container is None or container.is_linked_to_previous:
                continue

            text = "\n".join(p.text.strip() for p in container.paragraphs if p.text.strip())

            if text and text not in seen:
                seen.add(text)
                bucket.append(text)

    return {"headers": headers, "footers": footers}


def _extract_notes(doc, part_reltype_substring: str, xml_tag: str) -> list[str]:
    """Extract footnote or endnote body text."""
    notes_part = None

    for rel in doc.part.rels.values():
        if part_reltype_substring in rel.reltype:
            notes_part = rel.target_part
            break

    if notes_part is None:
        return []

    texts = []
    element = getattr(notes_part, "element", None)

    if element is None:
        element = parse_xml(notes_part.blob)

    for note in element.findall(qn(xml_tag)):
        note_id = note.get(qn("w:id"))

        if note_id in ("-1", "0"):
            continue

        note_text = "".join(t.text for t in note.iter(qn("w:t")) if t.text).strip()

        if note_text:
            texts.append(note_text)

    return texts


def _extract_document_properties(doc) -> dict:
    """Pull core document properties (title/author/dates) into metadata."""
    props = doc.core_properties
    meta: dict = {}

    if props.title:
        meta["doc_title"] = props.title

    if props.author:
        meta["doc_author"] = props.author

    if props.subject:
        meta["doc_subject"] = props.subject

    if props.created:
        meta["doc_created"] = props.created.isoformat()

    if props.modified:
        meta["doc_modified"] = props.modified.isoformat()

    return meta


def _extract_image_captions(doc) -> list[str]:
    """Extract alt text/descriptions from inline images in DOCX."""
    captions = []

    for p in doc.paragraphs:
        for drawing in p._element.findall(f".//{qn('w:drawing')}"):
            for docPr in drawing.findall(f".//{qn('wp:docPr')}"):
                descr = docPr.get(qn("wp:descr")) or docPr.get("descr", "")

                if descr and descr.strip():
                    captions.append(f"[image: {descr.strip()}]")

        for obj in p._element.findall(f".//{qn('w:object')}"):
            ole = obj.find(f".//{qn('o:OLEObject')}")

            if ole is not None:
                prog_id = ole.get(qn("o:ProgID")) or "OLE object"
                captions.append(f"[embedded: {prog_id}]")

        for _smart_art in p._element.findall(f".//{qn('dgm:relents')}"):
            captions.append("[diagram: SmartArt]")

    return captions


def docx_table_to_markdown(table) -> str:
    """Convert a python-docx table to markdown table format."""
    rows = []

    for row in table.rows:
        cells = [markdown_cell(cell.text) for cell in row.cells]
        rows.append(cells)

    return render_table(rows)


# --- parse_docx_sections helpers (extracted to reduce C901) -----------------


@dataclass
class DocxSectionState:
    """Heading and numbering state for the document body walk."""

    sections: list[Section]
    heading_stack: list[tuple[int, str]]
    current_heading: str | None = None
    current_lines: list[str] = field(default_factory=list)
    numberer: ListNumberer = field(default_factory=lambda: ListNumberer({}))


def _detect_heading_level(p) -> int | None:
    """Detect heading level from paragraph style. Returns 0 for Title, 1-N for Heading N, None otherwise."""
    style_name = (p.style.name or "").lower() if p.style else ""

    if style_name == "title":
        return 0

    if style_name.startswith("heading"):
        suffix = style_name.replace("heading", "").strip()
        return int(suffix) if suffix.isdigit() else 1

    return None


def _apply_heading(state: DocxSectionState, heading_level: int, text: str) -> None:
    """Flush current content and update heading stack."""
    content = "\n".join(state.current_lines).strip()

    if content:
        state.sections.append(Section(state.current_heading, content))

    while state.heading_stack and state.heading_stack[-1][0] >= heading_level:
        state.heading_stack.pop()

    state.heading_stack.append((heading_level, text))
    state.current_heading = " > ".join(h for _, h in state.heading_stack)
    state.current_lines = []


def flush_table(state: DocxSectionState, md_table: str) -> None:
    """Flush preceding text and append the rendered table."""
    if md_table:
        content = "\n".join(state.current_lines).strip()

        if content:
            state.sections.append(Section(state.current_heading, content))

        state.sections.append(Section(state.current_heading, md_table, PageContentType.TABLE))
        state.current_lines = []


def process_section_paragraph(state: DocxSectionState, paragraph, text: str) -> None:
    text = text.strip()

    if not text:
        return

    heading_level = _detect_heading_level(paragraph)

    if heading_level is not None:
        _apply_heading(state, heading_level, text)
    else:
        list_info = _paragraph_list_info(paragraph)

        if list_info:
            num_id, level = list_info
            text = state.numberer.prefix(num_id, level) + text

        state.current_lines.append(text)


def _collect_docx_metadata(
    doc,
    page_numbers: list[int],
    paragraph_count: int,
    table_count: int,
    image_count: int,
) -> dict:
    """Build metadata dict from document statistics and properties."""
    metadata: dict = {}

    if page_numbers:
        metadata["page_start"] = page_numbers[0]
        metadata["page_end"] = page_numbers[-1]
        metadata["pages"] = sorted(set(page_numbers))

    metadata["paragraph_count"] = paragraph_count
    metadata["table_count"] = table_count
    metadata["image_count"] = image_count

    headers_footers = _extract_headers_footers(doc)

    if headers_footers["headers"]:
        metadata["header_text"] = "\n".join(headers_footers["headers"])

    if headers_footers["footers"]:
        metadata["footer_text"] = "\n".join(headers_footers["footers"])

    metadata.update(_extract_document_properties(doc))
    return metadata


# --- Public API -------------------------------------------------------------


def append_docx_paragraph(
    paragraph,
    text: str,
    numberer: ListNumberer,
    parts: list[str],
    page_numbers: list[int],
    current_page: int,
) -> int:
    if _has_page_break(paragraph):
        current_page += 1

    if text.strip():
        page_numbers.append(current_page)
        list_info = _paragraph_list_info(paragraph)

        if list_info:
            num_id, level = list_info
            text = numberer.prefix(num_id, level) + text

        parts.append(text)

    return current_page


def read_docx(file_path: Path) -> tuple[str, list[Section], dict]:
    """Read body order, sections and document metadata in one document pass."""
    validate_docx_archive(file_path)
    doc = docx.Document(str(file_path))
    formats = _get_numbering_formats(doc)
    numberer = ListNumberer(formats)
    parts: list[str] = []
    page_numbers: list[int] = []
    current_page = 1
    paragraph_count = 0
    table_count = 0

    body = doc.element.body
    para_tag = qn("w:p")
    tbl_tag = qn("w:tbl")

    para_by_element = {id(p._element): p for p in doc.paragraphs}
    tbl_by_element = {id(t._element): t for t in doc.tables}

    state = DocxSectionState(sections=[], heading_stack=[], numberer=ListNumberer(formats))

    for child in body:
        if child.tag == para_tag:
            paragraph_count += 1
            p = para_by_element.get(id(child))

            if p is not None:
                text = _paragraph_full_text(p)
                process_section_paragraph(state, p, text)
                current_page = append_docx_paragraph(p, text, numberer, parts, page_numbers, current_page)
        elif child.tag == tbl_tag:
            table_count += 1
            table = tbl_by_element.get(id(child))

            if table is not None:
                markdown = docx_table_to_markdown(table)
                flush_table(state, markdown)

                if markdown:
                    parts.extend(["", markdown])

    image_count = sum(1 for _ in body.iter(qn("w:drawing")))

    captions = _extract_image_captions(doc)
    parts.extend(captions)

    footnotes = _extract_notes(doc, "footnotes", "w:footnote")

    for i, note in enumerate(footnotes, start=1):
        parts.append(f"[footnote {i}: {note}]")

    endnotes = _extract_notes(doc, "endnotes", "w:endnote")

    for i, note in enumerate(endnotes, start=1):
        parts.append(f"[endnote {i}: {note}]")

    metadata = _collect_docx_metadata(doc, page_numbers, paragraph_count, table_count, image_count)

    content = "\n".join(state.current_lines).strip()

    if content:
        state.sections.append(Section(state.current_heading, content))

    note_lines = [f"[footnote {i}: {n}]" for i, n in enumerate(footnotes, start=1)]
    note_lines += [f"[endnote {i}: {n}]" for i, n in enumerate(endnotes, start=1)]

    if note_lines:
        state.sections.append(Section("Footnotes", "\n".join(note_lines)))

    text = "\n".join(parts)
    return text, state.sections or [Section(None, text)], metadata


def parse_docx(file_path: Path) -> tuple[str, dict]:
    """Parse DOCX and return flat text and metadata."""
    text, _sections, metadata = read_docx(file_path)
    return text, metadata


def parse_docx_sections(file_path: Path) -> list[Section]:
    """Split DOCX by Heading/Title styles, preserving table kinds."""
    _text, sections, _metadata = read_docx(file_path)
    return sections
